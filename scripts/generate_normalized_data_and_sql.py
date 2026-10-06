import openpyxl
import json
import re
import os

def generate_normalized_data_and_sql():
    excel_path = 'JADWAL KBM - PIKET TA.2026-2027.xlsx'
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws_kbm = wb['JADWAL KBM']

    # 1. DATA GURU TERNORMALISASI (31 Guru + 1 Kepsek + 3 Staf Piket)
    # Nama dirapikan: Title Case, gelar EYD/PUEBI baku, NIP & kode guru
    teachers_master = {
        1: {
            'kode': 1,
            'nama': 'Eti Anisyah, S.Pd.',
            'nip': None,
            'jabatan': 'Guru IPS',
            'email': 'eti.anisyah@smpn8karawangbarat.sch.id',
            'role': 'GURU'
        },
        2: {
            'kode': 2,
            'nama': 'Mardiyah, M.Pd.',
            'nip': '19720725 200501 2 007',
            'jabatan': 'PKS Kurikulum / Guru Matematika',
            'email': 'mardiyah@smpn8karawangbarat.sch.id',
            'role': 'GURU'
        },
        3: {
            'kode': 3,
            'nama': 'Mimin Suherman, M.Pd.',
            'nip': None,
            'jabatan': 'Guru PJOK',
            'email': 'mimin.suherman@smpn8karawangbarat.sch.id',
            'role': 'GURU'
        },
        4: {
            'kode': 4,
            'nama': 'Siti Barkah, S.Pd.',
            'nip': None,
            'jabatan': 'Guru Matematika',
            'email': 'siti.barkah@smpn8karawangbarat.sch.id',
            'role': 'GURU'
        },
        5: {
            'kode': 5,
            'nama': 'Santi Susanti, S.Pd.',
            'nip': None,
            'jabatan': 'Guru Bahasa Inggris',
            'email': 'santi.susanti@smpn8karawangbarat.sch.id',
            'role': 'GURU'
        },
        6: {
            'kode': 6,
            'nama': 'Setiadi Purnomo, S.Pd.',
            'nip': None,
            'jabatan': 'Guru Bahasa Indonesia',
            'email': 'setiadi.purnomo@smpn8karawangbarat.sch.id',
            'role': 'GURU'
        },
        7: {
            'kode': 7,
            'nama': 'Ulfah Nurul Hikmawati, S.Pd.',
            'nip': None,
            'jabatan': 'Guru Bahasa Sunda',
            'email': 'ulfah.nurul@smpn8karawangbarat.sch.id',
            'role': 'GURU'
        },
        8: {
            'kode': 8,
            'nama': 'Renita Dean Sari, S.Pd.',
            'nip': None,
            'jabatan': 'Guru PPKn',
            'email': 'renita.dean@smpn8karawangbarat.sch.id',
            'role': 'GURU'
        },
        9: {
            'kode': 9,
            'nama': 'Niken Norma Yunita, S.Pd.',
            'nip': None,
            'jabatan': 'Guru PJOK',
            'email': 'niken.norma@smpn8karawangbarat.sch.id',
            'role': 'GURU'
        },
        10: {
            'kode': 10,
            'nama': 'Ahmad Husen Multiana, S.Pd.',
            'nip': None,
            'jabatan': 'Guru Bahasa Inggris',
            'email': 'ahmad.husen@smpn8karawangbarat.sch.id',
            'role': 'GURU'
        },
        11: {
            'kode': 11,
            'nama': 'Desty Nurbaety, S.Pd.',
            'nip': None,
            'jabatan': 'Guru IPA',
            'email': 'desty.nurbaety@smpn8karawangbarat.sch.id',
            'role': 'GURU'
        },
        12: {
            'kode': 12,
            'nama': 'Wahyudin, S.E., S.Pd.I',
            'nip': None,
            'jabatan': 'Guru Seni Budaya (SBK)',
            'email': 'wahyudin@smpn8karawangbarat.sch.id',
            'role': 'GURU'
        },
        13: {
            'kode': 13,
            'nama': 'Mimin Aminah, S.Ag.',
            'nip': None,
            'jabatan': 'Guru PAI',
            'email': 'mimin.aminah@smpn8karawangbarat.sch.id',
            'role': 'GURU'
        },
        14: {
            'kode': 14,
            'nama': 'Elva Mariana, S.Pd.',
            'nip': None,
            'jabatan': 'Guru Bahasa Inggris',
            'email': 'elva.mariana@smpn8karawangbarat.sch.id',
            'role': 'GURU'
        },
        15: {
            'kode': 15,
            'nama': 'Dede Sumarna, S.Kom.',
            'nip': None,
            'jabatan': 'Guru Informatika (TIK)',
            'email': 'dede.sumarna@smpn8karawangbarat.sch.id',
            'role': 'GURU'
        },
        16: {
            'kode': 16,
            'nama': 'Umaya Habibah, S.E.',
            'nip': None,
            'jabatan': 'Guru IPS',
            'email': 'umaya.habibah@smpn8karawangbarat.sch.id',
            'role': 'GURU'
        },
        17: {
            'kode': 17,
            'nama': 'Siti Mustikah, S.Pd.',
            'nip': None,
            'jabatan': 'Guru PPKn',
            'email': 'siti.mustikah@smpn8karawangbarat.sch.id',
            'role': 'GURU'
        },
        18: {
            'kode': 18,
            'nama': 'Awaliatush Sholihah, S.Pd.',
            'nip': None,
            'jabatan': 'Guru Matematika',
            'email': 'awaliatush.sholihah@smpn8karawangbarat.sch.id',
            'role': 'GURU'
        },
        19: {
            'kode': 19,
            'nama': 'Khossol Jawad, S.Pd.',
            'nip': None,
            'jabatan': 'Guru IPA',
            'email': 'khossol.jawad@smpn8karawangbarat.sch.id',
            'role': 'GURU'
        },
        20: {
            'kode': 20,
            'nama': 'Rosemalyna Khatimah, S.Pd.',
            'nip': None,
            'jabatan': 'Guru Bahasa Indonesia',
            'email': 'rosemalyna.khatimah@smpn8karawangbarat.sch.id',
            'role': 'GURU'
        },
        21: {
            'kode': 21,
            'nama': 'Nova Indriana, S.Pd.',
            'nip': None,
            'jabatan': 'Guru Bahasa Indonesia',
            'email': 'nova.indriana@smpn8karawangbarat.sch.id',
            'role': 'GURU'
        },
        22: {
            'kode': 22,
            'nama': 'Sri Dewi Hardayani, S.Pd.',
            'nip': None,
            'jabatan': 'Guru SBK & Bahasa Sunda',
            'email': 'sri.dewi@smpn8karawangbarat.sch.id',
            'role': 'GURU'
        },
        23: {
            'kode': 23,
            'nama': 'Maya Hardini, S.Pd.',
            'nip': None,
            'jabatan': 'Guru Bahasa Indonesia',
            'email': 'maya.hardini@smpn8karawangbarat.sch.id',
            'role': 'GURU'
        },
        24: {
            'kode': 24,
            'nama': 'Mirani Kartini, S.Pd.',
            'nip': None,
            'jabatan': 'Guru Matematika',
            'email': 'mirani.kartini@smpn8karawangbarat.sch.id',
            'role': 'GURU'
        },
        25: {
            'kode': 25,
            'nama': 'Siti Nurjamilah, S.Pd.I',
            'nip': None,
            'jabatan': 'Guru PAI',
            'email': 'siti.nurjamilah@smpn8karawangbarat.sch.id',
            'role': 'GURU'
        },
        26: {
            'kode': 26,
            'nama': 'Damanhuri, S.Pd.I., M.Pd.',
            'nip': None,
            'jabatan': 'Guru PAI & TIK',
            'email': 'damanhuri@smpn8karawangbarat.sch.id',
            'role': 'GURU'
        },
        27: {
            'kode': 27,
            'nama': 'Siska Nurnianti, S.Pd.',
            'nip': None,
            'jabatan': 'Guru Bahasa Inggris & SBK',
            'email': 'siska.nurnianti@smpn8karawangbarat.sch.id',
            'role': 'GURU'
        },
        28: {
            'kode': 28,
            'nama': 'Iwan Irnawan, S.Pd.',
            'nip': None,
            'jabatan': 'Guru IPA',
            'email': 'iwan.irnawan@smpn8karawangbarat.sch.id',
            'role': 'GURU'
        },
        29: {
            'kode': 29,
            'nama': 'Dede Supriyanto, S.Pd.',
            'nip': None,
            'jabatan': 'Guru IPA',
            'email': 'dede.supriyanto@smpn8karawangbarat.sch.id',
            'role': 'GURU'
        },
        30: {
            'kode': 30,
            'nama': 'Qurotul Aini, S.Pd., M.Pd.',
            'nip': None,
            'jabatan': 'Guru IPS',
            'email': 'qurotul.aini@smpn8karawangbarat.sch.id',
            'role': 'GURU'
        },
        31: {
            'kode': 31,
            'nama': 'Muhamad Nurseha, S.Pd.',
            'nip': None,
            'jabatan': 'Guru Bimbingan Konseling (BK)',
            'email': 'muhamad.nurseha@smpn8karawangbarat.sch.id',
            'role': 'GURU'
        },
    }

    # Pimpinan & Staf Tambahan
    extra_staff = [
        {
            'kode': 99,
            'nama': 'Mamay Abdullah, S.Pd., M.Pd.',
            'nip': '19700724 199802 1 003',
            'jabatan': 'Kepala Sekolah',
            'email': 'kepsek@smpn8karawangbarat.sch.id',
            'role': 'ADMIN'
        },
        {
            'kode': 101,
            'nama': 'Ana',
            'nip': None,
            'jabatan': 'Tenaga Kependidikan / Petugas Piket',
            'email': 'ana@smpn8karawangbarat.sch.id',
            'role': 'GURU'
        },
        {
            'kode': 102,
            'nama': 'Imi Suminar',
            'nip': None,
            'jabatan': 'Tenaga Kependidikan / Petugas Piket',
            'email': 'imi.suminar@smpn8karawangbarat.sch.id',
            'role': 'GURU'
        },
        {
            'kode': 103,
            'nama': 'Hanni Apnianti',
            'nip': None,
            'jabatan': 'Tenaga Kependidikan / Petugas Piket',
            'email': 'hanni.apnianti@smpn8karawangbarat.sch.id',
            'role': 'GURU'
        },
    ]

    # 2. MATA PELAJARAN (12 Mapel)
    subjects_master = {
        'A': {'kode': 'A', 'kode_mapel': 'PAIBP-A', 'nama_mapel': 'Pendidikan Agama Islam & Budi Pekerti', 'alias': 'PAIBP'},
        'B': {'kode': 'B', 'kode_mapel': 'PPKN-B', 'nama_mapel': 'Pendidikan Pancasila & Kewarganegaraan', 'alias': 'PPKn'},
        'C': {'kode': 'C', 'kode_mapel': 'BIND-C', 'nama_mapel': 'Bahasa Indonesia', 'alias': 'B. Indonesia'},
        'D': {'kode': 'D', 'kode_mapel': 'MTK-D', 'nama_mapel': 'Matematika', 'alias': 'Matematika'},
        'E': {'kode': 'E', 'kode_mapel': 'IPA-E', 'nama_mapel': 'Ilmu Pengetahuan Alam', 'alias': 'IPA'},
        'F': {'kode': 'F', 'kode_mapel': 'IPS-F', 'nama_mapel': 'Ilmu Pengetahuan Sosial', 'alias': 'IPS'},
        'G': {'kode': 'G', 'kode_mapel': 'BING-G', 'nama_mapel': 'Bahasa Inggris', 'alias': 'B. Inggris'},
        'H': {'kode': 'H', 'kode_mapel': 'SBK-H', 'nama_mapel': 'Seni Budaya & Keterampilan', 'alias': 'SBK'},
        'I': {'kode': 'I', 'kode_mapel': 'PJOK-I', 'nama_mapel': 'Pendidikan Jasmani, Olahraga, & Kesehatan', 'alias': 'PJOK'},
        'J': {'kode': 'J', 'kode_mapel': 'TIK-J', 'nama_mapel': 'Teknologi Informasi & Komunikasi / Informatika', 'alias': 'TIK'},
        'K': {'kode': 'K', 'kode_mapel': 'BSUN-K', 'nama_mapel': 'Bahasa Sunda', 'alias': 'B. Sunda'},
        'L': {'kode': 'L', 'kode_mapel': 'BK-L', 'nama_mapel': 'Bimbingan & Konseling', 'alias': 'BK'},
    }

    # 3. KELAS / RUANGAN (23 Rombel dengan Wali Kelas)
    classes_list = [ws_kbm.cell(6, c).value for c in range(4, 27)]
    wali_codes = [ws_kbm.cell(66, c).value for c in range(4, 27)]
    
    classes_master = []
    for cls, w_code in zip(classes_list, wali_codes):
        tingkat = 'VII' if cls.startswith('7') else ('VIII' if cls.startswith('8') else 'IX')
        wali_info = teachers_master.get(int(w_code)) if w_code else None
        classes_master.append({
            'nama_kelas': cls,
            'nama_ruangan': f'Ruang Kelas {cls}',
            'tingkat': tingkat,
            'wali_kelas_kode': int(w_code) if w_code else None,
            'wali_kelas_nama': wali_info['nama'] if wali_info else None,
            'kode_qr': f'QR-RUANG-KLS{cls}'
        })

    # 4. JADWAL PIKET
    picket_master = [
        {'hari': 'Senin', 'petugas': 'Ahmad Husen Multiana, S.Pd.', 'kode_guru': 10, 'catatan': 'Petugas Piket Hari Senin'},
        {'hari': 'Selasa', 'petugas': 'Ana', 'kode_guru': 101, 'catatan': 'Petugas Piket Hari Selasa'},
        {'hari': 'Rabu', 'petugas': 'Imi Suminar', 'kode_guru': 102, 'catatan': 'Petugas Piket Hari Rabu'},
        {'hari': 'Kamis', 'petugas': 'Hanni Apnianti', 'kode_guru': 103, 'catatan': 'Petugas Piket Hari Kamis'},
        {'hari': 'Jumat', 'petugas': 'Damanhuri, S.Pd.I., M.Pd.', 'kode_guru': 26, 'catatan': 'Petugas Piket Hari Jumat'},
    ]

    # 5. JADWAL KBM (Sesi Mengajar)
    row_time_map = {
        8: (1, '07:30:00', '08:10:00'),
        9: (2, '08:10:00', '08:50:00'),
        10: (3, '08:50:00', '09:30:00'),
        12: (4, '09:50:00', '10:30:00'),
        13: (5, '10:30:00', '11:10:00'),
        14: (6, '11:10:00', '11:50:00'),
        16: (7, '12:30:00', '13:10:00'),
        17: (8, '13:10:00', '13:50:00'),
        
        20: (1, '07:30:00', '08:10:00'),
        21: (2, '08:10:00', '08:50:00'),
        22: (3, '08:50:00', '09:30:00'),
        24: (4, '09:50:00', '10:30:00'),
        25: (5, '10:30:00', '11:10:00'),
        26: (6, '11:10:00', '11:50:00'),
        28: (7, '12:30:00', '13:10:00'),
        29: (8, '13:10:00', '13:50:00'),

        32: (1, '07:00:00', '07:40:00'),
        33: (2, '07:40:00', '08:20:00'),
        34: (3, '08:20:00', '09:00:00'),
        35: (4, '09:00:00', '09:40:00'),
        37: (5, '10:00:00', '10:40:00'),
        38: (6, '10:40:00', '11:20:00'),
        39: (7, '11:20:00', '12:00:00'),
        41: (8, '12:40:00', '13:20:00'),
        42: (9, '13:20:00', '14:00:00'),

        45: (1, '07:00:00', '07:40:00'),
        46: (2, '07:40:00', '08:20:00'),
        47: (3, '08:20:00', '09:00:00'),
        48: (4, '09:00:00', '09:40:00'),
        50: (5, '10:00:00', '10:40:00'),
        51: (6, '10:40:00', '11:20:00'),
        52: (7, '11:20:00', '12:00:00'),
        54: (8, '12:40:00', '13:20:00'),
        55: (9, '13:20:00', '14:00:00'),

        58: (1, '07:30:00', '08:00:00'),
        59: (2, '08:00:00', '08:30:00'),
        60: (3, '08:30:00', '09:00:00'),
        61: (4, '09:00:00', '09:30:00'),
        63: (5, '10:00:00', '10:30:00'),
        64: (6, '10:30:00', '11:00:00'),
    }

    day_rows = {
        'Senin': [8, 9, 10, 12, 13, 14, 16, 17],
        'Selasa': [20, 21, 22, 24, 25, 26, 28, 29],
        'Rabu': [32, 33, 34, 35, 37, 38, 39, 41, 42],
        'Kamis': [45, 46, 47, 48, 50, 51, 52, 54, 55],
        'Jumat': [58, 59, 60, 61, 63, 64],
    }

    schedules_master = []

    for day, rows in day_rows.items():
        for c_idx, cls in enumerate(classes_list, 4):
            current_block = None
            for r in rows:
                val = str(ws_kbm.cell(r, c_idx).value or '').strip()
                m = re.match(r'^(\d+)([A-Za-z]?)$', val)
                t_code = int(m.group(1)) if m else None
                s_code = m.group(2).upper() if (m and m.group(2)) else ('L' if t_code == 31 else None)
                jam, start_t, end_t = row_time_map[r]
                
                if current_block and current_block['teacher_code'] == t_code and current_block['subject_code'] == s_code:
                    current_block['jam_selesai_ke'] = jam
                    current_block['jam_selesai'] = end_t
                else:
                    if current_block and current_block['teacher_code'] is not None:
                        schedules_master.append(current_block)
                    current_block = {
                        'kelas': cls,
                        'hari': day,
                        'teacher_code': t_code,
                        'subject_code': s_code,
                        'jam_mulai_ke': jam,
                        'jam_selesai_ke': jam,
                        'jam_mulai': start_t,
                        'jam_selesai': end_t,
                    }
            if current_block and current_block['teacher_code'] is not None:
                schedules_master.append(current_block)

    # Jam BK (Layanan Bimbingan Konseling)
    bk_rows = [
        (18, 'Senin', '13:50:00', '14:30:00'),
        (30, 'Selasa', '13:50:00', '14:30:00'),
        (43, 'Rabu', '14:00:00', '14:40:00'),
        (56, 'Kamis', '14:00:00', '14:40:00'),
        (65, 'Jumat', '11:00:00', '11:40:00'),
    ]
    for r, day, s_time, e_time in bk_rows:
        for c_idx, cls in enumerate(classes_list, 4):
            val = str(ws_kbm.cell(r, c_idx).value or '').strip()
            if '31' in val:
                schedules_master.append({
                    'kelas': cls,
                    'hari': day,
                    'teacher_code': 31,
                    'subject_code': 'L',
                    'jam_mulai_ke': 0,
                    'jam_selesai_ke': 0,
                    'jam_mulai': s_time,
                    'jam_selesai': e_time,
                })

    # Tambahkan label jam_ke (e.g. '1-2', '3-5', 'BK')
    for s in schedules_master:
        if s['jam_mulai_ke'] == 0:
            s['jam_ke'] = 'BK'
        elif s['jam_mulai_ke'] == s['jam_selesai_ke']:
            s['jam_ke'] = f"{s['jam_mulai_ke']}"
        else:
            s['jam_ke'] = f"{s['jam_mulai_ke']}-{s['jam_selesai_ke']}"

    # Simpan JSON
    full_data = {
        'sekolah': {
            'nama': 'SMP Negeri 8 Karawang Barat',
            'tahun_pelajaran': '2026/2027',
            'semester': 'Ganjil',
            'kepala_sekolah': {
                'nama': 'Mamay Abdullah, S.Pd., M.Pd.',
                'nip': '19700724 199802 1 003'
            },
            'pks_kurikulum': {
                'nama': 'Mardiyah, M.Pd.',
                'nip': '19720725 200501 2 007'
            }
        },
        'teachers': list(teachers_master.values()),
        'extra_staff': extra_staff,
        'subjects': list(subjects_master.values()),
        'classes': classes_master,
        'picket_schedule': picket_master,
        'schedules': schedules_master,
    }

    os.makedirs('scripts', exist_ok=True)
    with open('scripts/data_kbm_dan_piket_normalized.json', 'w', encoding='utf-8') as f:
        json.dump(full_data, f, indent=2, ensure_ascii=False)
    print("Berhasil membuat scripts/data_kbm_dan_piket_normalized.json")

    # 6. GENERATE SQL SEED SCRIPT
    sql_lines = []
    sql_lines.append("-- ====================================================================")
    sql_lines.append("-- MIGRASI & SEED MASTER DATA KBM & PIKET SMP NEGERI 8 KARAWANG BARAT")
    sql_lines.append("-- TAHUN PELAJARAN 2026 / 2027 (SEMESTER GANJIL)")
    sql_lines.append("-- Jalankan script ini pada Supabase SQL Editor")
    sql_lines.append("-- ====================================================================\n")
    
    sql_lines.append("-- 1. PENYESUAIAN SKEMA TABEL (Agar Lebih Proper & Siap untuk KBM & Piket)")
    sql_lines.append("CREATE EXTENSION IF NOT EXISTS pgcrypto;")
    sql_lines.append("CREATE EXTENSION IF NOT EXISTS \"uuid-ossp\";\n")
    
    sql_lines.append("-- Tambah kolom kode_guru pada profiles jika belum ada")
    sql_lines.append("ALTER TABLE public.profiles ADD COLUMN IF NOT EXISTS kode_guru INT;")
    sql_lines.append("CREATE UNIQUE INDEX IF NOT EXISTS idx_profiles_kode_guru ON public.profiles(kode_guru) WHERE kode_guru IS NOT NULL;\n")

    sql_lines.append("-- Tambah kolom tingkat & wali_kelas_id pada rooms jika belum ada")
    sql_lines.append("ALTER TABLE public.rooms ADD COLUMN IF NOT EXISTS tingkat VARCHAR(10);")
    sql_lines.append("ALTER TABLE public.rooms ADD COLUMN IF NOT EXISTS wali_kelas_id UUID REFERENCES public.profiles(id) ON DELETE SET NULL;\n")

    sql_lines.append("-- Tambah kolom jam_ke pada schedules jika belum ada")
    sql_lines.append("ALTER TABLE public.schedules ADD COLUMN IF NOT EXISTS jam_ke VARCHAR(20);\n")

    sql_lines.append("-- Tambah tabel picket_schedules (Jadwal Guru Piket)")
    sql_lines.append("""CREATE TABLE IF NOT EXISTS public.picket_schedules (
  id UUID PRIMARY KEY DEFAULT gen_random_uuid(),
  hari VARCHAR(20) NOT NULL, -- 'Senin', 'Selasa', 'Rabu', 'Kamis', 'Jumat'
  teacher_id UUID REFERENCES public.profiles(id) ON DELETE CASCADE,
  nama_petugas VARCHAR(150) NOT NULL,
  catatan TEXT,
  created_at TIMESTAMPTZ DEFAULT NOW(),
  CONSTRAINT unique_picket_per_day UNIQUE (hari, teacher_id)
);

ALTER TABLE public.picket_schedules ENABLE ROW LEVEL SECURITY;

DO $$ BEGIN
  CREATE POLICY "Allow read picket schedules" ON public.picket_schedules FOR SELECT TO authenticated USING (true);
EXCEPTION WHEN duplicate_object THEN null;
END $$;

DO $$ BEGIN
  CREATE POLICY "Allow admin manage picket schedules" ON public.picket_schedules FOR ALL TO authenticated USING (public.is_admin());
EXCEPTION WHEN duplicate_object THEN null;
END $$;
""")

    sql_lines.append("\n-- 2. SEED SEMUA AKUN GURU & PIMPINAN KE auth.users & public.profiles")
    sql_lines.append("DO $$")
    sql_lines.append("DECLARE")
    sql_lines.append("  new_id UUID;")
    sql_lines.append("BEGIN")

    all_users = list(teachers_master.values()) + extra_staff
    for u in all_users:
        nama_esc = u['nama'].replace("'", "''")
        jabatan_esc = u['jabatan'].replace("'", "''")
        nip_val = f"'{u['nip']}'" if u['nip'] else "NULL"
        email = u['email']
        role = u['role']
        kode = u['kode']

        sql_lines.append(f"""
  -- Akun: {u['nama']} (Kode: {kode})
  SELECT id INTO new_id FROM auth.users WHERE email = '{email}';
  IF new_id IS NULL THEN
    new_id := gen_random_uuid();
    INSERT INTO auth.users (
      id, instance_id, email, encrypted_password, email_confirmed_at,
      raw_app_meta_data, raw_user_meta_data, created_at, updated_at, role, aud
    ) VALUES (
      new_id, '00000000-0000-0000-0000-000000000000', '{email}',
      crypt('guru123', gen_salt('bf')), NOW(),
      '{{\"provider\":\"email\",\"providers\":[\"email\"]}}',
      jsonb_build_object('role', '{role}', 'nama', '{nama_esc}', 'nip', {nip_val}, 'jabatan', '{jabatan_esc}', 'kode_guru', {kode}),
      NOW(), NOW(), 'authenticated', 'authenticated'
    );
  END IF;

  INSERT INTO public.profiles (id, email, nama, nip, role, jabatan, kode_guru)
  VALUES (new_id, '{email}', '{nama_esc}', {nip_val}, '{role}'::public.user_role, '{jabatan_esc}', {kode})
  ON CONFLICT (id) DO UPDATE SET
    nama = EXCLUDED.nama,
    nip = EXCLUDED.nip,
    role = EXCLUDED.role,
    jabatan = EXCLUDED.jabatan,
    kode_guru = EXCLUDED.kode_guru;
""")

    sql_lines.append("END $$;\n")

    # 3. SEED MATA PELAJARAN (Tahan terhadap constraint nama_mapel maupun kode_mapel)
    sql_lines.append("-- 3. SEED 12 MATA PELAJARAN")
    sql_lines.append("DO $$")
    sql_lines.append("DECLARE")
    sql_lines.append("  sub RECORD;")
    sql_lines.append("  target_sub_id UUID;")
    sql_lines.append("BEGIN")
    sql_lines.append("  FOR sub IN")
    sql_lines.append("    SELECT * FROM (VALUES")
    sub_values = []
    for s in subjects_master.values():
        n = s['nama_mapel'].replace("'", "''")
        k = s['kode_mapel']
        sub_values.append(f"      ('{n}', '{k}')")
    sql_lines.append(",\n".join(sub_values))
    sql_lines.append("    ) AS t(nama, kode)")
    sql_lines.append("  LOOP")
    sql_lines.append("    -- Cari id mapel yang sudah ada berdasarkan kode atau nama")
    sql_lines.append("    SELECT id INTO target_sub_id FROM public.subjects WHERE kode_mapel = sub.kode LIMIT 1;")
    sql_lines.append("    IF target_sub_id IS NULL THEN")
    sql_lines.append("      SELECT id INTO target_sub_id FROM public.subjects WHERE nama_mapel = sub.nama LIMIT 1;")
    sql_lines.append("    END IF;")
    sql_lines.append("    IF target_sub_id IS NOT NULL THEN")
    sql_lines.append("      UPDATE public.subjects SET nama_mapel = sub.nama, kode_mapel = sub.kode WHERE id = target_sub_id;")
    sql_lines.append("    ELSE")
    sql_lines.append("      INSERT INTO public.subjects (nama_mapel, kode_mapel) VALUES (sub.nama, sub.kode);")
    sql_lines.append("    END IF;")
    sql_lines.append("  END LOOP;")
    sql_lines.append("END $$;\n")

    # 4. SEED 23 RUANG KELAS / ROMBEL & PENETAPAN WALI KELAS
    sql_lines.append("-- 4. SEED 23 RUANG KELAS / ROMBEL & PENETAPAN WALI KELAS")
    sql_lines.append("DO $$")
    sql_lines.append("DECLARE")
    sql_lines.append("  rm RECORD;")
    sql_lines.append("  target_rm_id UUID;")
    sql_lines.append("  wali_id UUID;")
    sql_lines.append("BEGIN")
    for c in classes_master:
        cls_name = c['nama_ruangan'].replace("'", "''")
        tingkat = c['tingkat']
        kode_qr = c['kode_qr']
        w_kode = c['wali_kelas_kode']
        sql_lines.append(f"""
  -- Ruang: {cls_name}
  SELECT id INTO wali_id FROM public.profiles WHERE kode_guru = {w_kode} LIMIT 1;
  SELECT id INTO target_rm_id FROM public.rooms WHERE nama_ruangan = '{cls_name}' LIMIT 1;
  IF target_rm_id IS NULL THEN
    SELECT id INTO target_rm_id FROM public.rooms WHERE kode_qr = '{kode_qr}' LIMIT 1;
  END IF;

  IF target_rm_id IS NOT NULL THEN
    UPDATE public.rooms SET
      nama_ruangan = '{cls_name}',
      gedung = 'Gedung Kelas {tingkat}',
      deskripsi = 'Ruang Belajar {cls_name} TP 2026/2027',
      tingkat = '{tingkat}',
      kode_qr = '{kode_qr}',
      wali_kelas_id = wali_id
    WHERE id = target_rm_id;
  ELSE
    INSERT INTO public.rooms (nama_ruangan, gedung, deskripsi, tingkat, kode_qr, wali_kelas_id)
    VALUES ('{cls_name}', 'Gedung Kelas {tingkat}', 'Ruang Belajar {cls_name} TP 2026/2027', '{tingkat}', '{kode_qr}', wali_id);
  END IF;
""")
    sql_lines.append("END $$;\n")

    # 5. SEED JADWAL PIKET
    sql_lines.append("-- 5. SEED JADWAL PIKET GURU")
    sql_lines.append("DELETE FROM public.picket_schedules;")
    for p in picket_master:
        hari = p['hari']
        kode_g = p['kode_guru']
        nama_pet = p['petugas'].replace("'", "''")
        catatan = p['catatan'].replace("'", "''")
        sql_lines.append(f"""
INSERT INTO public.picket_schedules (hari, teacher_id, nama_petugas, catatan)
VALUES (
  '{hari}',
  (SELECT id FROM public.profiles WHERE kode_guru = {kode_g} LIMIT 1),
  '{nama_pet}',
  '{catatan}'
);""")

    # 6. SEED JADWAL KBM
    sql_lines.append("\n-- 6. SEED SELURUH JADWAL PELAJARAN KBM (391 Sesi)")
    sql_lines.append("-- Membersihkan jadwal lama untuk sinkronisasi bersih:")
    sql_lines.append("DELETE FROM public.schedules;\n")

    sql_lines.append("INSERT INTO public.schedules (teacher_id, subject_id, room_id, hari, jam_mulai, jam_selesai, jam_ke)")
    sql_lines.append("VALUES")

    value_rows = []
    for s in schedules_master:
        cls = s['kelas']
        t_code = s['teacher_code']
        s_code = s['subject_code']
        hari = s['hari']
        j_mulai = s['jam_mulai']
        j_selesai = s['jam_selesai']
        j_ke = s['jam_ke']

        s_obj = subjects_master.get(s_code)
        kode_mapel = s_obj['kode_mapel'] if s_obj else f"{s_code}"
        nama_mapel_esc = s_obj['nama_mapel'].replace("'", "''") if s_obj else ""

        val = f"""  (
    (SELECT id FROM public.profiles WHERE kode_guru = {t_code} LIMIT 1),
    COALESCE((SELECT id FROM public.subjects WHERE kode_mapel = '{kode_mapel}' LIMIT 1), (SELECT id FROM public.subjects WHERE nama_mapel = '{nama_mapel_esc}' LIMIT 1)),
    COALESCE((SELECT id FROM public.rooms WHERE nama_ruangan = 'Ruang Kelas {cls}' LIMIT 1), (SELECT id FROM public.rooms WHERE kode_qr = 'QR-RUANG-KLS{cls}' LIMIT 1)),
    '{hari}',
    '{j_mulai}',
    '{j_selesai}',
    '{j_ke}'
  )"""
        value_rows.append(val)

    sql_lines.append(",\n".join(value_rows) + ";\n")

    sql_lines.append("-- 7. VERIFIKASI SEED DATA")
    sql_lines.append("SELECT 'Total Guru' AS entitas, COUNT(*) AS jumlah FROM public.profiles WHERE role = 'GURU'")
    sql_lines.append("UNION ALL SELECT 'Total Ruang Kelas', COUNT(*) FROM public.rooms")
    sql_lines.append("UNION ALL SELECT 'Total Mata Pelajaran', COUNT(*) FROM public.subjects")
    sql_lines.append("UNION ALL SELECT 'Total Jadwal KBM', COUNT(*) FROM public.schedules")
    sql_lines.append("UNION ALL SELECT 'Total Jadwal Piket', COUNT(*) FROM public.picket_schedules;")

    sql_content = "\n".join(sql_lines)
    with open('supabase/seed_kbm_dan_piket.sql', 'w', encoding='utf-8') as f:
        f.write(sql_content)
    print("Berhasil membuat supabase/seed_kbm_dan_piket.sql")

if __name__ == '__main__':
    generate_normalized_data_and_sql()
